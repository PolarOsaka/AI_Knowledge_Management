import sys
from pathlib import Path
sys.path.append(str(Path(__file__).resolve().parent.parent))
import hashlib
import json
import traceback
import fitz  # PyMuPDF
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from configuration import ROOT_DIR, TREE_PATH, DOCUMENTS_DIR, OUTPUT_DIR, INDEX_PATH, ERROR_PATH, SUPPORTED_EXTENSIONS, CHUNK_SIZE, CHUNK_OVERLAP

# ============================================================
# 1. UTILITIES
# ============================================================

def stable_id(value):
    """Create a deterministic identifier from a string."""
    return hashlib.sha256(
        value.encode("utf-8")
    ).hexdigest()[:20]


def clean_text(value):
    """Normalize whitespace in extracted text."""
    if value is None:
        return ""

    return " ".join(str(value).split()).strip()


def chunk_text(text, chunk_size=CHUNK_SIZE,
               overlap=CHUNK_OVERLAP):
    """Split long text into overlapping character chunks."""
    text = clean_text(text)

    if not text:
        return []

    if overlap < 0 or overlap >= chunk_size:
        raise ValueError(
            "Overlap must be >= 0 and smaller than chunk size."
        )

    chunks = []
    start = 0

    while start < len(text):
        end = min(start + chunk_size, len(text))

        # Prefer ending at a word boundary when possible.
        if end < len(text):
            boundary = text.rfind(" ", start, end)
            if boundary > start + chunk_size // 2:
                end = boundary

        piece = text[start:end].strip()

        if piece:
            chunks.append(piece)

        if end >= len(text):
            break

        start = max(end - overlap, start + 1)

    return chunks


def read_json(path):
    with open(path, "r", encoding="utf-8") as file:
        return json.load(file)


# ============================================================
# 3. LOAD PROJECT METADATA FROM KNOWLEDGE TREE
# ============================================================

def collect_projects(node, projects):
    """Recursively collect project nodes from the tree."""
    if node.get("node_type") == "project":
        code = node.get("project_code")

        if code:
            if code in projects:
                raise ValueError(
                    f"Duplicate project code in tree: {code}"
                )

            projects[code] = {
                "project_code": code,
                **node.get("metadata", {}),
            }

    for child in node.get("children", {}).values():
        collect_projects(child, projects)


if not TREE_PATH.exists():
    raise FileNotFoundError(
        f"Knowledge tree not found: {TREE_PATH}"
    )

if not DOCUMENTS_DIR.exists():
    raise FileNotFoundError(
        f"Documents directory not found: {DOCUMENTS_DIR}\n"
        "Create data/documents/ and add project folders."
    )

tree = read_json(TREE_PATH)
projects_by_code = {}
collect_projects(tree, projects_by_code)

print(f"Projects in knowledge tree: {len(projects_by_code)}")


# ============================================================
# 4. DOCUMENT TEXT EXTRACTORS
# Each result contains text + a source locator.
# ============================================================

def extract_pdf(path):
    units = []

    with fitz.open(path) as pdf:
        for page_index, page in enumerate(pdf):
            text = page.get_text("text")

            if clean_text(text):
                units.append({
                    "locator": {
                        "type": "pdf_page",
                        "page": page_index + 1,
                    },
                    "text": text,
                })

    return units


def extract_pptx(path):
    units = []
    presentation = Presentation(path)

    for slide_index, slide in enumerate(
        presentation.slides, start=1
    ):
        text_parts = []

        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                if shape.text.strip():
                    text_parts.append(shape.text)

            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [
                        cell.text.strip()
                        for cell in row.cells
                    ]
                    text_parts.append(" | ".join(cells))

        if text_parts:
            units.append({
                "locator": {
                    "type": "pptx_slide",
                    "slide": slide_index,
                },
                "text": "\n".join(text_parts),
            })

    return units


def extract_docx(path):
    units = []
    document = Document(path)

    # Paragraphs: paragraph number is more stable than page
    # because Word pagination depends on rendering and layout.
    for index, paragraph in enumerate(
        document.paragraphs, start=1
    ):
        text = paragraph.text.strip()

        if text:
            units.append({
                "locator": {
                    "type": "docx_paragraph",
                    "paragraph": index,
                    "style": paragraph.style.name
                    if paragraph.style else None,
                },
                "text": text,
            })

    # Tables are indexed separately by table and row.
    for table_index, table in enumerate(
        document.tables, start=1
    ):
        for row_index, row in enumerate(
            table.rows, start=1
        ):
            cells = [
                cell.text.strip()
                for cell in row.cells
            ]
            text = " | ".join(
                value for value in cells if value
            )

            if text:
                units.append({
                    "locator": {
                        "type": "docx_table_row",
                        "table": table_index,
                        "row": row_index,
                    },
                    "text": text,
                })

    return units


def extract_xlsx(path):
    units = []

    # data_only=True reads cached formula results when available.
    # It does not calculate Excel formulas.
    workbook = load_workbook(
        path,
        read_only=True,
        data_only=True,
    )

    try:
        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows()

            header_row = next(rows, None)

            if header_row is None:
                continue

            headers = []
            for col_index, cell in enumerate(
                header_row, start=1
            ):
                header = clean_text(cell.value)
                headers.append(
                    header or f"Column_{col_index}"
                )

            # Excel row number starts at 2 after the header.
            for row_number, row in enumerate(rows, start=2):
                fields = []

                for col_index, cell in enumerate(row):
                    value = clean_text(cell.value)

                    if not value:
                        continue

                    header = (
                        headers[col_index]
                        if col_index < len(headers)
                        else f"Column_{col_index + 1}"
                    )

                    fields.append(f"{header}: {value}")

                text = " | ".join(fields)

                if text:
                    units.append({
                        "locator": {
                            "type": "xlsx_row",
                            "sheet": worksheet.title,
                            "row": row_number,
                        },
                        "text": text,
                    })

    finally:
        workbook.close()

    return units


EXTRACTORS = {
    ".pdf": extract_pdf,
    ".pptx": extract_pptx,
    ".docx": extract_docx,
    ".xlsx": extract_xlsx,
}


# ============================================================
# 5. BUILD DOCUMENT INDEX
# One JSON object per text chunk.
# ============================================================

index_records = []
errors = []
seen_document_ids = set()

files = sorted(
    path for path in DOCUMENTS_DIR.rglob("*")
    if path.is_file()
    and path.suffix.lower() in SUPPORTED_EXTENSIONS
    and not path.name.startswith("~$")
    and not path.name.startswith(".")
)

print(f"Supported files discovered: {len(files)}")

for path in files:
    relative_path = path.relative_to(DOCUMENTS_DIR)

    # Expected path: data/documents/<Project Code>/file.ext
    if len(relative_path.parts) < 2:
        errors.append({
            "file": str(relative_path),
            "error": (
                "File is not inside a project-code folder."
            ),
        })
        continue

    project_code = relative_path.parts[0]

    if project_code not in projects_by_code:
        errors.append({
            "file": str(relative_path),
            "error": (
                f"Project code {project_code!r} "
                "does not exist in the knowledge tree."
            ),
        })
        continue

    project_metadata = projects_by_code[project_code]

    relative_source_path = path.relative_to(
        ROOT_DIR
    ).as_posix()

    document_id = stable_id(relative_source_path)

    if document_id in seen_document_ids:
        errors.append({
            "file": str(relative_path),
            "error": "Duplicate document identifier.",
        })
        continue

    seen_document_ids.add(document_id)

    try:
        extractor = EXTRACTORS[path.suffix.lower()]
        units = extractor(path)

        file_chunk_count = 0

        for unit_index, unit in enumerate(units):
            chunks = chunk_text(unit["text"])

            for chunk_index, chunk in enumerate(chunks):
                chunk_id = stable_id(
                    f"{document_id}|{unit_index}|{chunk_index}"
                )

                record = {
                    "chunk_id": chunk_id,
                    "document_id": document_id,
                    "project_code": project_code,
                    "project_name":
                        project_metadata.get("project_name"),
                    "project_type":
                        project_metadata.get("project_type"),
                    "practice_area":
                        project_metadata.get("practice_area"),
                    "start_year":
                        project_metadata.get("start_year"),
                    "client_name":
                        project_metadata.get("client_name"),
                    "status":
                        project_metadata.get("status"),
                    "sharepoint_folder":
                        project_metadata.get("sharepoint_folder"),
                    "file_name": path.name,
                    "file_type": path.suffix.lower(),
                    "source_path": relative_source_path,
                    "source_locator": unit["locator"],
                    "chunk_index": chunk_index,
                    "text": chunk,
                }

                index_records.append(record)
                file_chunk_count += 1

        print(
            f"[OK] {relative_source_path} "
            f"-> {file_chunk_count} chunks"
        )

        if not units:
            print(
                f"[WARN] No extractable text: "
                f"{relative_source_path}"
            )

    except Exception as exc:
        errors.append({
            "file": str(relative_path),
            "error": str(exc),
            "traceback": traceback.format_exc(),
        })

        print(f"[ERROR] {relative_source_path}: {exc}")


# ============================================================
# 6. WRITE OUTPUTS
# ============================================================

with open(INDEX_PATH, "w", encoding="utf-8") as file:
    for record in index_records:
        file.write(
            json.dumps(record, ensure_ascii=False) + "\n"
        )

with open(ERROR_PATH, "w", encoding="utf-8") as file:
    json.dump(errors, file, indent=2, ensure_ascii=False)


# ============================================================
# 7. SUMMARY
# ============================================================

print("\n========== INDEXING SUMMARY ==========")
print(f"Projects in tree : {len(projects_by_code)}")
print(f"Files discovered : {len(files)}")
print(f"Chunks indexed   : {len(index_records)}")
print(f"Errors           : {len(errors)}")
print(f"Index output     : {INDEX_PATH}")
print(f"Error report     : {ERROR_PATH}")
